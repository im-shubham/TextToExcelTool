import concurrent.futures
import datetime as dt
import io
import multiprocessing
import uuid

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.config import Settings
from app.main import MSG_EMPTY, MSG_SAVE_FAILED, MSG_SUCCESS, MSG_UNAVAILABLE, create_app
from app.workbook import SHEET_NAME, last_used_row, parse_blocks

from .compare import first_difference, sheet_signature, style_signature


def post(client, date, shift, activities, sid=None, confirm=False):
    return client.post("/api/submissions", json={
        "submission_id": sid or str(uuid.uuid4()),
        "date": date,
        "shift": shift,
        "activities": activities,
        "confirm_existing": confirm,
    })


def master(settings: Settings):
    return openpyxl.load_workbook(settings.master_path)


def rows_for(settings, date):
    ws = master(settings)[SHEET_NAME]
    out = []
    for b in parse_blocks(ws):
        if b.date == dt.date.fromisoformat(date):
            out.extend((b.shift, ws.cell(r, 4).value) for r in b.activity_rows)
    return out


# ------------------------------------------------------------------ basic appends


def test_single_activity_updates_master(client, settings):
    r = post(client, "2026-10-24", "B", ["Main line scanner holders fixed"])
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "saved" and body["added"] == 1 and body["message"] == MSG_SUCCESS
    assert rows_for(settings, "2026-10-24") == [("B", "Main line scanner holders fixed")]


def test_multiple_activities_blank_lines_ignored_wording_kept(client, settings):
    acts = ["Main line scanner holders fixed", "", "   ", "  Packing line ESD chairs arranged  ",
            "Post line workstation setup completed"]
    r = post(client, "2026-10-24", "B", acts)
    assert r.json()["added"] == 3
    assert [a for _, a in rows_for(settings, "2026-10-24")] == [
        "Main line scanner holders fixed", "Packing line ESD chairs arranged",
        "Post line workstation setup completed"]


def test_new_block_layout_and_fixed_values(client, settings):
    before = last_used_row(master(settings)[SHEET_NAME])
    post(client, "2026-10-24", "A", ["one", "two", "three"])
    ws = master(settings)[SHEET_NAME]
    h = before + 2
    assert ws.cell(h - 1, 1).value is None  # one blank separator row
    assert [ws.cell(h, c).value for c in range(1, 9)] == [
        "Sl.no", "Date", "Shift", "Work activity", "Assigned CDL person", "Work status", "Shift DRI", "Remarks"]
    assert [ws.cell(h + i, 1).value for i in range(1, 8)] == [1, 2, 3, 4, 5, 6, 7]
    assert all(ws.cell(h + i, 2).value == dt.datetime(2026, 10, 24) for i in range(1, 8))
    assert ws.cell(h + 1, 3).value == "A"
    assert ws.cell(h + 1, 5).value == settings.cdl_text
    assert ws.cell(h + 1, 6).value == "Completed"
    assert ws.cell(h + 1, 7).value == settings.shift_dri
    assert ws.cell(h + 1, 8).value == "CLOSED"
    merges = {str(m) for m in ws.merged_cells.ranges}
    for col in "CEFGH":
        assert f"{col}{h + 1}:{col}{h + 7}" in merges


def test_new_block_copies_existing_formatting(client, settings):
    ws0 = master(settings)[SHEET_NAME]
    ref = [b for b in parse_blocks(ws0) if b.last_row - b.first_row == 6][-1]
    post(client, "2026-10-24", "A", ["x"])
    ws = master(settings)[SHEET_NAME]
    new = parse_blocks(ws)[-1]
    for off in range(8):
        for c in range(1, 9):
            assert style_signature(ws.cell(new.header_row + off, c)) == \
                style_signature(ws0.cell(ref.header_row + off, c)), (off, c)


def test_next_day_preserves_previous_day_and_history(client, settings):
    original = master(settings)
    hist_rows = last_used_row(original[SHEET_NAME])
    post(client, "2026-10-24", "B", ["day one a", "day one b"])
    post(client, "2026-10-25", "A", ["day two a"])
    wb = master(settings)
    assert rows_for(settings, "2026-10-24") == [("B", "day one a"), ("B", "day one b")]
    assert rows_for(settings, "2026-10-25") == [("A", "day two a")]
    # Historical rows: values, styles, merges, widths and heights untouched.
    before, after = sheet_signature(original[SHEET_NAME], hist_rows), sheet_signature(wb[SHEET_NAME], hist_rows)
    assert after == before, first_difference(before, after)
    # Other worksheets untouched.
    for name in original.sheetnames:
        if name != SHEET_NAME:
            assert sheet_signature(wb[name]) == sheet_signature(original[name])
    assert wb.sheetnames == original.sheetnames


def test_more_than_seven_activities_grow_the_block(client, settings):
    acts = [f"activity {i}" for i in range(12)]
    post(client, "2026-10-24", "A", acts)
    ws = master(settings)[SHEET_NAME]
    b = parse_blocks(ws)[-1]
    assert b.last_row - b.first_row + 1 == 12
    assert [ws.cell(r, 4).value for r in b.activity_rows] == acts
    assert f"E{b.first_row}:E{b.last_row}" in {str(m) for m in ws.merged_cells.ranges}


def test_formula_like_text_is_stored_as_text(client, settings):
    post(client, "2026-10-24", "A", ["=HYPERLINK(\"http://x\")", "+1-2"])
    wb = master(settings)
    ws = wb[SHEET_NAME]
    b = parse_blocks(ws)[-1]
    assert ws.cell(b.first_row, 4).data_type == "s"
    assert ws.cell(b.first_row, 4).value == "=HYPERLINK(\"http://x\")"


# ------------------------------------------------------------------ existing date handling


def test_existing_date_requires_confirmation(client, settings):
    post(client, "2026-10-24", "B", ["first"])
    snapshot = settings.master_path.read_bytes()
    r = post(client, "2026-10-24", "B", ["second"])
    assert r.status_code == 409
    assert r.json()["code"] == "date_exists"
    assert r.json()["entries"] == [{"shift": "B", "activity_count": 1}]
    assert settings.master_path.read_bytes() == snapshot  # nothing written

    r = post(client, "2026-10-24", "B", ["second", "third"], confirm=True)
    assert r.status_code == 200 and r.json()["added"] == 2
    assert [a for _, a in rows_for(settings, "2026-10-24")] == ["first", "second", "third"]
    # Filled into the same block rather than a new one.
    assert len([b for b in parse_blocks(master(settings)[SHEET_NAME])
                if b.date == dt.date(2026, 10, 24)]) == 1


def test_adding_beyond_block_capacity_continues_numbering(client, settings):
    post(client, "2026-10-24", "B", [f"a{i}" for i in range(5)])
    post(client, "2026-10-24", "B", [f"b{i}" for i in range(4)], confirm=True)
    ws = master(settings)[SHEET_NAME]
    blocks = [b for b in parse_blocks(ws) if b.date == dt.date(2026, 10, 24)]
    assert len(blocks) == 2
    assert ws.cell(blocks[1].first_row, 1).value == 8
    assert [a for _, a in rows_for(settings, "2026-10-24")] == \
        [f"a{i}" for i in range(5)] + [f"b{i}" for i in range(4)]


def test_both_shifts_same_date(client, settings):
    post(client, "2026-10-24", "A", ["morning work"])
    r = post(client, "2026-10-24", "B", ["evening work"], confirm=True)
    assert r.status_code == 200
    assert sorted(rows_for(settings, "2026-10-24")) == [("A", "morning work"), ("B", "evening work")]
    assert client.get("/api/dates/2026-10-24").json()["entries"] == [
        {"shift": "A", "activity_count": 1}, {"shift": "B", "activity_count": 1}]


def test_existing_historical_date_detected(client, settings):
    ws = master(settings)[SHEET_NAME]
    hist = next(b for b in parse_blocks(ws) if b.date)
    r = post(client, hist.date.isoformat(), "A", ["late entry"])
    assert r.status_code == 409


# ------------------------------------------------------------------ idempotency


def test_retrying_same_request_does_not_duplicate(client, settings):
    sid = str(uuid.uuid4())
    acts = ["only once"]
    first = post(client, "2026-10-24", "A", acts, sid=sid)
    again = post(client, "2026-10-24", "A", acts, sid=sid)
    again_confirmed = post(client, "2026-10-24", "A", acts, sid=sid, confirm=True)
    assert first.json()["status"] == "saved"
    assert again.status_code == 200 and again.json()["status"] == "duplicate"
    assert again_confirmed.json()["status"] == "duplicate"
    assert rows_for(settings, "2026-10-24") == [("A", "only once")]


def test_concurrent_double_click_saves_once(client, settings):
    sid = str(uuid.uuid4())
    with concurrent.futures.ThreadPoolExecutor(8) as ex:
        results = list(ex.map(lambda _: post(client, "2026-10-24", "A", ["click"], sid=sid), range(8)))
    statuses = sorted(r.json()["status"] for r in results)
    assert statuses.count("saved") == 1 and statuses.count("duplicate") == 7
    assert rows_for(settings, "2026-10-24") == [("A", "click")]


def test_reused_id_with_different_content_is_rejected(client, settings):
    sid = str(uuid.uuid4())
    post(client, "2026-10-24", "A", ["one"], sid=sid)
    r = post(client, "2026-10-24", "A", ["two"], sid=sid, confirm=True)
    assert r.status_code == 409 and r.json()["code"] == "submission_conflict"
    assert rows_for(settings, "2026-10-24") == [("A", "one")]


# ------------------------------------------------------------------ concurrency


def test_concurrent_submissions_no_lost_updates(client, settings):
    dates = [(dt.date(2027, 3, 1) + dt.timedelta(days=i)).isoformat() for i in range(12)]
    with concurrent.futures.ThreadPoolExecutor(12) as ex:
        results = list(ex.map(lambda d: post(client, d, "A", [f"work on {d}", "second"]), dates))
    assert all(r.status_code == 200 for r in results), [r.text for r in results]
    for d in dates:
        assert rows_for(settings, d) == [("A", f"work on {d}"), ("A", "second")]


def _submit_in_process(args):
    settings, date = args
    with TestClient(create_app(settings)) as c:
        return post(c, date, "B", [f"process {date}"]).status_code


def test_concurrent_processes_no_lost_updates(settings, client):
    dates = [(dt.date(2027, 4, 1) + dt.timedelta(days=i)).isoformat() for i in range(6)]
    with multiprocessing.get_context("spawn").Pool(6) as pool:
        codes = pool.map(_submit_in_process, [(settings, d) for d in dates])
    assert codes == [200] * 6
    for d in dates:
        assert rows_for(settings, d) == [("B", f"process {d}")]


# ------------------------------------------------------------------ persistence


def test_master_survives_restart_and_is_not_reseeded(settings):
    with TestClient(create_app(settings)) as c:
        post(c, "2026-10-24", "B", ["before restart"])
    with TestClient(create_app(settings)) as c:  # a brand-new app instance = restart/redeploy
        assert rows_for(settings, "2026-10-24") == [("B", "before restart")]
        post(c, "2026-10-25", "A", ["after restart"])
    assert rows_for(settings, "2026-10-24") == [("B", "before restart")]
    assert rows_for(settings, "2026-10-25") == [("A", "after restart")]


def test_missing_master_after_init_is_reported_not_recreated(client, settings):
    settings.master_path.unlink()
    r = post(client, "2026-10-24", "A", ["x"])
    assert r.status_code == 503 and r.json()["message"] == MSG_UNAVAILABLE
    with TestClient(create_app(settings)):
        pass
    assert not settings.master_path.exists()  # seed is not silently copied over lost data


def test_missing_sheet_is_reported_without_replacement(client, settings):
    wb = master(settings)
    wb[SHEET_NAME].title = "Renamed"
    wb.save(settings.master_path)
    before = settings.master_path.read_bytes()
    r = post(client, "2026-10-24", "A", ["x"])
    assert r.status_code == 500 and r.json()["code"] == "sheet_missing"
    assert settings.master_path.read_bytes() == before


def test_save_failure_is_reported_and_master_untouched(client, settings, monkeypatch):
    before = settings.master_path.read_bytes()

    def boom(self, *a, **k):
        raise OSError("disk full")

    monkeypatch.setattr(openpyxl.Workbook, "save", boom)
    r = post(client, "2026-10-24", "A", ["x"])
    assert r.status_code == 500 and r.json()["message"] == MSG_SAVE_FAILED
    assert settings.master_path.read_bytes() == before
    assert not list(settings.data_dir.glob(".tmp-*"))


def test_backup_kept_before_each_save(client, settings):
    post(client, "2026-10-24", "A", ["x"])
    assert len(list((settings.data_dir / "backups").glob("*.xlsx"))) == 1


# ------------------------------------------------------------------ validation


def test_empty_submission_rejected(client, settings):
    before = settings.master_path.read_bytes()
    r = post(client, "2026-10-24", "A", ["", "   "])
    assert r.status_code == 422 and r.json()["message"] == MSG_EMPTY
    assert settings.master_path.read_bytes() == before


@pytest.mark.parametrize("field,value", [("shift", "C"), ("date", "2026-13-01"), ("submission_id", "nope")])
def test_invalid_input_rejected(client, field, value):
    body = {"submission_id": str(uuid.uuid4()), "date": "2026-10-24", "shift": "A", "activities": ["x"]}
    body[field] = value
    r = client.post("/api/submissions", json=body)
    assert r.status_code == 422


def test_too_long_activity_rejected(client, settings):
    r = post(client, "2026-10-24", "A", ["x" * (settings.max_activity_length + 1)])
    assert r.status_code == 422 and r.json()["code"] == "too_long"


# ------------------------------------------------------------------ download


def test_download_is_latest_snapshot_and_does_not_modify(client, settings):
    post(client, "2026-10-24", "A", ["downloaded"])
    before = settings.master_path.read_bytes()
    r = client.get("/api/workbook/download")
    assert r.status_code == 200
    assert r.content == before
    assert settings.master_path.read_bytes() == before
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert SHEET_NAME in wb.sheetnames
    assert "attachment" in r.headers["content-disposition"]


def test_status_reports_latest_entry(client):
    post(client, "2026-10-24", "B", ["x"])
    s = client.get("/api/workbook/status").json()
    assert s["latest_entry"] == {"date": "2026-10-24", "shift": "B"}
