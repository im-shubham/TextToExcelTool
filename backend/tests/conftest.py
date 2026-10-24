import datetime as dt
import os
import shutil
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.config import DEFAULT_CDL_PERSONS, DEFAULT_SHIFT_DRI, Settings
from app.main import create_app

# Set REAL_WORKBOOK=/path/to/GOKUL 26-27.xlsx to run every test against the real workbook as well.
REAL_WORKBOOK = os.environ.get("REAL_WORKBOOK")


def build_synthetic_workbook(path: Path) -> None:
    """A small workbook with the same layout as the real one (blocks, merges, styles, other sheets)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Work activity"
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="FFBDD7EE")
    headers = ["Sl.no", "Date", "Shift", "Work activity", "Assigned CDL person", "Work status", "Shift DRI", "Remarks"]

    def block(h, date, shift, acts, values=True):
        for c, t in enumerate(headers, 1):
            cell = ws.cell(h, c, t if values else None)
            cell.font, cell.fill, cell.border = header_font, header_fill, box
            cell.alignment = Alignment(horizontal="center")
        for i in range(7):
            r = h + 1 + i
            for c in range(1, 9):
                ws.cell(r, c).border = box
            ws.cell(r, 4).alignment = Alignment(horizontal="center", wrap_text=True)
            ws.cell(r, 2).number_format = "mm-dd-yy"
            if values:
                ws.cell(r, 1, i + 1)
                ws.cell(r, 2, dt.datetime(date.year, date.month, date.day))
                if i < len(acts):
                    ws.cell(r, 4, acts[i])
        for c in (3, 5, 6, 7, 8):
            ws.merge_cells(start_row=h + 1, end_row=h + 7, start_column=c, end_column=c)
        if values:
            ws.cell(h + 1, 3, shift)
            ws.cell(h + 1, 5, "\n".join(DEFAULT_CDL_PERSONS))
            ws.cell(h + 1, 6, "Completed")
            ws.cell(h + 1, 7, DEFAULT_SHIFT_DRI)
            ws.cell(h + 1, 8, "CLOSED")

    block(1, dt.date(2026, 8, 20), "A", ["first activity", "second activity"])
    block(10, dt.date(2026, 8, 21), "B ", ["third activity", "fourth", "fifth"])
    block(19, None, None, [], values=False)  # pre-formatted empty block typed ahead of time
    ws.column_dimensions["D"].width = 70

    other = wb.create_sheet("work 2026 -27")
    other.append(["Sl No", "Floor", "Work Description"])
    other.append([1, "TESS3 3F", "Clean booth extension"])
    other.merge_cells("A5:C5")
    esd = wb.create_sheet("ESD Chiar retrofit")
    esd.append(["s.no", "Date", "Line", "No of chairs", "Cost saved"])
    esd.append([1, "21/7/2026", "Rolex", 30, "=SUM(D2*300)"])
    wb.save(path)


@pytest.fixture(params=["synthetic"] + (["real"] if REAL_WORKBOOK else []))
def seed_workbook(request, tmp_path_factory) -> Path:
    d = tmp_path_factory.mktemp("seed")
    path = d / "seed.xlsx"
    if request.param == "real":
        shutil.copyfile(REAL_WORKBOOK, path)
    else:
        build_synthetic_workbook(path)
    return path


@pytest.fixture
def settings(tmp_path, seed_workbook) -> Settings:
    return Settings(data_dir=tmp_path / "data", master_name="GOKUL 26-27.xlsx", seed_path=seed_workbook,
                    lock_timeout=60)


@pytest.fixture
def client(settings):
    from fastapi.testclient import TestClient

    with TestClient(create_app(settings)) as c:
        yield c
