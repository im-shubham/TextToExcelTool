"""Reading and appending daily blocks on the "Work activity" sheet.

The sheet is a sequence of daily blocks separated by a blank row:

    Sl.no | Date | Shift | Work activity | Assigned CDL person | Work status | Shift DRI | Remarks   <- header
    1     | date | A     | activity 1    | (merged C,E,F,G,H over the data rows)                    <- data rows
    ...
    7     | date |       | (blank)                                                                  <- unused rows keep Sl.no/date
    <blank>

New blocks copy the cell styles of the most recent well-formed block so they are
indistinguishable from the ones typed by hand.
"""

from __future__ import annotations

import datetime as dt
from copy import copy
from dataclasses import dataclass, field

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

SHEET_NAME = "Work activity"
HEADERS = (
    "Sl.no",
    "Date",
    "Shift",
    "Work activity",
    "Assigned CDL person",
    "Work status",
    "Shift DRI",
    "Remarks",
)
NCOLS = len(HEADERS)
# Columns that are merged vertically across a block's data rows: Shift, CDL, status, DRI, remarks.
MERGED_COLS = (3, 5, 6, 7, 8)
COL_SERIAL, COL_DATE, COL_SHIFT, COL_ACTIVITY, COL_CDL, COL_STATUS, COL_DRI, COL_REMARKS = range(1, 9)
MIN_BLOCK_ROWS = 7
# Columns inspected when deciding whether a row is in use (a little wider than the table).
SCAN_COLS = 12


class DateAlreadyHasActivities(Exception):
    def __init__(self, entries: list[dict]):
        super().__init__("The selected date already has activities")
        self.entries = entries


@dataclass
class Block:
    header_row: int
    first_row: int
    last_row: int
    date: dt.date | None
    shift: str | None
    activity_rows: list[int] = field(default_factory=list)


@dataclass
class FixedValues:
    cdl: str
    status: str
    dri: str
    remarks: str


@dataclass
class AppendResult:
    rows: list[int]
    new_block_headers: list[int]
    extended_block_header: int | None


def clean_activity(text: str) -> str:
    """Trim surrounding whitespace and drop characters Excel cannot store; wording is otherwise untouched."""
    return ILLEGAL_CHARACTERS_RE.sub("", text).strip()


def _value(ws: Worksheet, row: int, col: int):
    # Read without ws.cell(), which would create empty cells as a side effect.
    cell = ws._cells.get((row, col))
    return cell.value if cell is not None else None


def _has_value(ws: Worksheet, row: int, col: int) -> bool:
    v = _value(ws, row, col)
    return v is not None and not (isinstance(v, str) and not v.strip())


def _row_has_value(ws: Worksheet, row: int, ncols: int = SCAN_COLS) -> bool:
    return any(_has_value(ws, row, c) for c in range(1, ncols + 1))


def _is_header(ws: Worksheet, row: int) -> bool:
    v = _value(ws, row, COL_SERIAL)
    return isinstance(v, str) and v.strip().lower() == "sl.no"


def _as_date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return None


def _merged_rows_spanning(ws: Worksheet) -> set[tuple[int, int]]:
    """(row, row+1) pairs that are joined by a vertical merge in the table columns."""
    pairs = set()
    for mr in ws.merged_cells.ranges:
        if mr.min_col > NCOLS:
            continue
        for r in range(mr.min_row, mr.max_row):
            pairs.add((r, r + 1))
    return pairs


def parse_blocks(ws: Worksheet) -> list[Block]:
    joined = _merged_rows_spanning(ws)
    max_row = ws.max_row
    blocks = []
    for h in range(1, max_row + 1):
        if not _is_header(ws, h):
            continue
        r = h + 1
        last = h
        # A block continues while rows carry data or are tied to the previous row by a merge.
        while r <= max_row and not _is_header(ws, r):
            if _row_has_value(ws, r, NCOLS) or (r - 1, r) in joined:
                last = r
                r += 1
            else:
                break
        if last == h:
            continue
        block = Block(header_row=h, first_row=h + 1, last_row=last, date=None, shift=None)
        for row in range(block.first_row, block.last_row + 1):
            if block.date is None:
                block.date = _as_date(_value(ws, row, COL_DATE))
            if block.shift is None:
                s = _value(ws, row, COL_SHIFT)
                if isinstance(s, str) and s.strip():
                    block.shift = s.strip().upper()
            if _has_value(ws, row, COL_ACTIVITY):
                block.activity_rows.append(row)
        blocks.append(block)
    return blocks


def entries_for_date(ws: Worksheet, date: dt.date) -> list[dict]:
    summary: dict[str, int] = {}
    for b in parse_blocks(ws):
        if b.date == date:
            key = b.shift or "?"
            summary[key] = summary.get(key, 0) + len(b.activity_rows)
    return [{"shift": s, "activity_count": n} for s, n in sorted(summary.items())]


def last_used_row(ws: Worksheet) -> int:
    for r in range(ws.max_row, 0, -1):
        if _row_has_value(ws, r):
            return r
    return 0


# ---------------------------------------------------------------- styles


@dataclass
class BlockStyles:
    header: list
    first: list
    middle: list
    last: list


def _fallback_styles() -> BlockStyles:
    """Used only if the sheet has no well-formed block to copy from."""
    from openpyxl.cell.cell import Cell
    from openpyxl.workbook import Workbook

    thin = Side(style="thin")
    tmp = Workbook().active

    def make(**kw):
        c = Cell(tmp)
        for k, v in kw.items():
            setattr(c, k, v)
        return c._style

    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    header = make(font=Font(bold=True), fill=PatternFill("solid", fgColor="FFD9E1F2"),
                  alignment=Alignment(horizontal="center"), border=box)
    body = [make(border=box, alignment=Alignment(horizontal="center", vertical="center", wrap_text=True))
            for _ in range(NCOLS)]
    body[COL_DATE - 1] = make(border=box, alignment=Alignment(horizontal="center"), number_format="mm-dd-yy")
    return BlockStyles([header] * NCOLS, body, body, body)


def reference_styles(ws: Worksheet, blocks: list[Block]) -> BlockStyles:
    merged_e = {(mr.min_row, mr.max_row) for mr in ws.merged_cells.ranges
                if mr.min_col == mr.max_col == COL_CDL}
    for b in reversed(blocks):
        if b.last_row - b.first_row + 1 == MIN_BLOCK_ROWS and (b.first_row, b.last_row) in merged_e and b.activity_rows:

            def row_styles(r):
                return [copy(ws.cell(r, c)._style) for c in range(1, NCOLS + 1)]
            return BlockStyles(
                header=row_styles(b.header_row),
                first=row_styles(b.first_row),
                middle=row_styles(b.first_row + 1),
                last=row_styles(b.last_row),
            )
    return _fallback_styles()


# ---------------------------------------------------------------- writing


def _set_text(cell, text: str) -> None:
    cell.value = text
    # Never let pasted text such as "=SUM(...)" become a formula.
    cell.data_type = "s"


def _clear_area(ws: Worksheet, top: int, bottom: int) -> None:
    """Remove empty, pre-formatted merges (e.g. a block template typed ahead of time) from rows top..bottom."""
    for mr in list(ws.merged_cells.ranges):
        if mr.max_row < top or mr.min_row > bottom or mr.min_col > SCAN_COLS:
            continue
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                if _has_value(ws, r, c):
                    raise RuntimeError(f"Refusing to overwrite existing data at row {r}")
        ws.unmerge_cells(str(mr))


def write_new_block(
    ws: Worksheet,
    header_row: int,
    date: dt.date,
    shift: str,
    activities: list[str],
    start_serial: int,
    styles: BlockStyles,
    fixed: FixedValues,
) -> list[int]:
    n = max(MIN_BLOCK_ROWS, len(activities))
    first, last = header_row + 1, header_row + n
    for r in range(header_row, last + 1):
        if _row_has_value(ws, r):
            raise RuntimeError(f"Refusing to overwrite existing data at row {r}")
    _clear_area(ws, header_row, last)

    for c in MERGED_COLS:
        ws.merge_cells(start_row=first, end_row=last, start_column=c, end_column=c)

    for c, title in enumerate(HEADERS, start=1):
        cell = ws.cell(header_row, c)
        cell._style = copy(styles.header[c - 1])
        cell.value = title

    when = dt.datetime(date.year, date.month, date.day)
    rows = []
    for i in range(n):
        r = first + i
        src = styles.first if i == 0 else styles.last if i == n - 1 else styles.middle
        for c in range(1, NCOLS + 1):
            ws.cell(r, c)._style = copy(src[c - 1])
        ws.cell(r, COL_SERIAL).value = start_serial + i
        ws.cell(r, COL_DATE).value = when
        if i < len(activities):
            _set_text(ws.cell(r, COL_ACTIVITY), activities[i])
            rows.append(r)

    ws.cell(first, COL_SHIFT).value = shift
    _set_text(ws.cell(first, COL_CDL), fixed.cdl)
    _set_text(ws.cell(first, COL_STATUS), fixed.status)
    _set_text(ws.cell(first, COL_DRI), fixed.dri)
    _set_text(ws.cell(first, COL_REMARKS), fixed.remarks)
    return rows


def _fill_free_rows(ws: Worksheet, block: Block, activities: list[str]) -> tuple[list[int], int]:
    """Put activities into the unused rows at the end of an existing block.

    Returns the rows written and the next serial number for any overflow.
    """
    after = max(block.activity_rows, default=block.first_row - 1)
    when = dt.datetime(block.date.year, block.date.month, block.date.day) if block.date else None
    written = []
    queue = list(activities)
    for r in range(after + 1, block.last_row + 1):
        if not queue:
            break
        d = ws.cell(r, COL_ACTIVITY)
        if isinstance(d, MergedCell) or _has_value(ws, r, COL_ACTIVITY):
            continue
        a = ws.cell(r, COL_SERIAL)
        prev = ws.cell(r - 1, COL_SERIAL).value
        if not isinstance(a, MergedCell) and a.value is None and isinstance(prev, int):
            a.value = prev + 1
        b = ws.cell(r, COL_DATE)
        if when and not isinstance(b, MergedCell) and b.value is None:
            b.value = when
        _set_text(d, queue.pop(0))
        written.append(r)
    serials = [ws.cell(r, COL_SERIAL).value for r in range(block.first_row, block.last_row + 1)]
    return written, max((v for v in serials if isinstance(v, int)), default=0) + 1


def append_activities(
    ws: Worksheet,
    date: dt.date,
    shift: str,
    activities: list[str],
    fixed: FixedValues,
    confirm_existing: bool,
) -> AppendResult:
    blocks = parse_blocks(ws)
    existing = [b for b in blocks if b.date == date]
    if existing and not confirm_existing:
        raise DateAlreadyHasActivities(entries_for_date(ws, date))

    styles = reference_styles(ws, blocks)
    result = AppendResult(rows=[], new_block_headers=[], extended_block_header=None)
    remaining = list(activities)
    start_serial = 1

    same_shift = [b for b in existing if b.shift == shift]
    if same_shift:
        target = same_shift[-1]
        written, start_serial = _fill_free_rows(ws, target, remaining)
        if written:
            result.rows.extend(written)
            result.extended_block_header = target.header_row
            remaining = remaining[len(written):]

    if remaining:
        used = last_used_row(ws)
        header_row = used + 2 if used else 1
        result.rows.extend(
            write_new_block(ws, header_row, date, shift, remaining, start_serial, styles, fixed)
        )
        result.new_block_headers.append(header_row)
    return result
