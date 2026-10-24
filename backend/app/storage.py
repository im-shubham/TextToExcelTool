"""Persistent storage of the single master workbook.

Every write goes through `MasterStore.transaction()`, which
  1. takes an exclusive OS-level file lock (safe across threads, worker processes and containers
     that share the volume),
  2. loads the latest saved workbook,
  3. lets the caller modify it,
  4. saves to a temporary file next to the master, re-opens it to verify the content,
  5. keeps a timestamped backup of the previous version and atomically renames the new file
     over the master, and
  6. releases the lock.

A crash at any point leaves either the old or the new master in place, never a partial file.
"""

from __future__ import annotations

import contextlib
import datetime as dt
import errno
import fcntl
import hashlib
import json
import os
import shutil
import time
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterator

import openpyxl
from openpyxl.packaging.custom import StringProperty
from openpyxl.workbook import Workbook

from .config import Settings
from .workbook import SHEET_NAME

# Processed submission ids are stored as custom document properties inside the workbook itself,
# so "activities appended" and "submission recorded" are saved in the same atomic write.
SUBMISSION_PROP_PREFIX = "DWA submission "
SUBMISSIONS_TO_REMEMBER = 500
INIT_MARKER = ".initialized"


class WorkbookUnavailable(Exception):
    pass


class SheetMissing(Exception):
    pass


class SaveFailed(Exception):
    pass


class LockTimeout(Exception):
    pass


def payload_hash(date: dt.date, shift: str, activities: list[str]) -> str:
    raw = json.dumps([date.isoformat(), shift, activities], ensure_ascii=False)
    return hashlib.sha256(raw.encode()).hexdigest()[:24]


def find_submission(wb: Workbook, submission_id: str) -> dict | None:
    name = SUBMISSION_PROP_PREFIX + submission_id
    for prop in wb.custom_doc_props.props:
        if prop.name == name:
            try:
                return json.loads(prop.value)
            except (TypeError, ValueError):
                return {}
    return None


def record_submission(wb: Workbook, submission_id: str, record: dict) -> None:
    props = wb.custom_doc_props
    ours = [p for p in props.props if p.name.startswith(SUBMISSION_PROP_PREFIX)]
    for old in ours[: max(0, len(ours) - SUBMISSIONS_TO_REMEMBER + 1)]:
        props.props.remove(old)
    value = json.dumps(record, separators=(",", ":"))
    props.append(StringProperty(name=SUBMISSION_PROP_PREFIX + submission_id, value=value[:255]))


@dataclass
class Transaction:
    wb: Workbook
    save: bool = True
    # Called with the freshly re-opened saved file; raise to reject the save.
    verify: Callable[[Workbook], None] = lambda wb: None


class MasterStore:
    def __init__(self, settings: Settings):
        self.settings = settings
        self.path = settings.master_path
        self.lock_path = settings.data_dir / ".master.lock"
        self.backup_dir = settings.data_dir / "backups"
        self.log_path = settings.data_dir / "submissions.log"

    # ------------------------------------------------------------ setup

    def initialize(self) -> None:
        """Copy the seed workbook into the data directory exactly once.

        After the first run a marker file exists; from then on a missing master is reported as an
        error instead of silently being replaced by the original template (which would lose data).
        """
        data_dir = self.settings.data_dir
        data_dir.mkdir(parents=True, exist_ok=True)
        marker = data_dir / INIT_MARKER
        if self.path.exists():
            marker.touch(exist_ok=True)
            return
        if marker.exists():
            return
        seed = self.settings.seed_path
        if seed and seed.exists():
            with self._locked():
                if not self.path.exists():
                    tmp = self.path.with_name(f".tmp-{uuid.uuid4().hex}-{self.path.name}")
                    shutil.copyfile(seed, tmp)
                    self._fsync_file(tmp)
                    os.replace(tmp, self.path)
                    self._fsync_dir()
                marker.touch(exist_ok=True)

    # ------------------------------------------------------------ locking

    @contextlib.contextmanager
    def _locked(self) -> Iterator[None]:
        self.settings.data_dir.mkdir(parents=True, exist_ok=True)
        fd = os.open(self.lock_path, os.O_RDWR | os.O_CREAT, 0o600)
        try:
            deadline = time.monotonic() + self.settings.lock_timeout
            while True:
                try:
                    fcntl.flock(fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
                    break
                except OSError as e:
                    if e.errno not in (errno.EAGAIN, errno.EACCES):
                        raise
                    if time.monotonic() > deadline:
                        raise LockTimeout("Timed out waiting for the master workbook lock")
                    time.sleep(0.05)
            try:
                yield
            finally:
                fcntl.flock(fd, fcntl.LOCK_UN)
        finally:
            os.close(fd)

    # ------------------------------------------------------------ io

    def _load(self) -> Workbook:
        if not self.path.is_file():
            raise WorkbookUnavailable(f"Master workbook not found at {self.path}")
        try:
            wb = openpyxl.load_workbook(self.path)
        except Exception as e:  # corrupted file, permission problem, ...
            raise WorkbookUnavailable(str(e)) from e
        if SHEET_NAME not in wb.sheetnames:
            raise SheetMissing(f'The master workbook has no "{SHEET_NAME}" worksheet')
        return wb

    def read(self) -> Workbook:
        """Load the latest saved workbook for read-only use (saves are atomic renames, so no lock needed)."""
        return self._load()

    def read_bytes(self) -> bytes:
        if not self.path.is_file():
            raise WorkbookUnavailable(f"Master workbook not found at {self.path}")
        return self.path.read_bytes()

    @staticmethod
    def _fsync_file(path: Path) -> None:
        with open(path, "rb") as f:
            os.fsync(f.fileno())

    def _fsync_dir(self) -> None:
        fd = os.open(self.settings.data_dir, os.O_RDONLY)
        try:
            os.fsync(fd)
        finally:
            os.close(fd)

    def _backup_current(self) -> None:
        self.backup_dir.mkdir(exist_ok=True)
        stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        shutil.copy2(self.path, self.backup_dir / f"{self.path.stem} {stamp}{self.path.suffix}")
        backups = sorted(self.backup_dir.glob(f"{self.path.stem} *{self.path.suffix}"))
        for old in backups[: max(0, len(backups) - self.settings.backups_to_keep)]:
            old.unlink(missing_ok=True)

    def _save(self, wb: Workbook, verify: Callable[[Workbook], None]) -> None:
        tmp = self.path.with_name(f".tmp-{uuid.uuid4().hex}-{self.path.name}")
        try:
            # Formulas on other sheets lose their cached values when re-saved; ask Excel to recalc.
            if wb.calculation is not None:
                wb.calculation.fullCalcOnLoad = True
            wb.save(tmp)
            self._fsync_file(tmp)
            check = openpyxl.load_workbook(tmp)
            if SHEET_NAME not in check.sheetnames:
                raise SaveFailed("Saved file is missing the activity sheet")
            verify(check)
            expected = hashlib.sha256(tmp.read_bytes()).hexdigest()
            self._backup_current()
            os.replace(tmp, self.path)
            self._fsync_dir()
            if hashlib.sha256(self.path.read_bytes()).hexdigest() != expected:
                raise SaveFailed("Master workbook content does not match what was written")
        except SaveFailed:
            raise
        except Exception as e:
            raise SaveFailed(str(e)) from e
        finally:
            tmp.unlink(missing_ok=True)

    @contextlib.contextmanager
    def transaction(self) -> Iterator["Transaction"]:
        """Exclusive load-modify-save. If the body raises or sets `save = False`, nothing is written."""
        with self._locked():
            tx = Transaction(self._load())
            yield tx
            if tx.save:
                self._save(tx.wb, tx.verify)

    def log(self, entry: dict) -> None:
        """Append-only audit trail (not used for correctness)."""
        with contextlib.suppress(OSError):
            with open(self.log_path, "a", encoding="utf-8") as f:
                f.write(json.dumps(entry, ensure_ascii=False) + "\n")
